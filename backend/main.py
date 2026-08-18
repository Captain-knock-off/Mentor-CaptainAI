from __future__ import annotations

import base64
import io
import os
import threading
from collections import defaultdict
from pathlib import Path
from typing import Any

import requests
from dotenv import load_dotenv
from fastapi import FastAPI, HTTPException, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
from pptx import Presentation

load_dotenv()

GROQ_API_KEY = os.getenv("GROQ_API_KEY")
TEXT_MODEL = os.getenv("TEXT_MODEL", "openai/gpt-oss-120b")
VISION_MODEL = os.getenv("VISION_MODEL", "qwen/qwen3.6-27b")
GROQ_URL = "https://api.groq.com/openai/v1/chat/completions"
MAX_FILE_SIZE = 10 * 1024 * 1024
MAX_FILES_PER_SESSION = 10
MAX_IMAGES_PER_REQUEST = 3
MAX_HISTORY_MESSAGES = 24
MAX_DOCUMENT_CHARS = 120_000

ALLOWED_EXTENSIONS = {"txt","md","markdown","csv","json","py","js","ts","html","css","xml","yaml","yml","log","ini","toml","sql","java","c","cpp","h","hpp","jsx","tsx","sh","bat","ps1","env","rtf","pdf","docx","pptx","xlsx","png","jpg","jpeg","webp","gif"}

SYSTEM_PROMPT = """
You are Mentor.CaptainAI, a modern AI tutor for students.

Your job is to help the student UNDERSTAND, not merely dump an answer.

TEACHING:
- Explain naturally and clearly.
- Match the student's likely level.
- Start simple, then add depth when useful.
- Break difficult problems into logical steps.
- Explain what formulas mean and define symbols/units.
- Use small examples when they improve understanding.
- Prefer teaching over unexplained final answers.
- Never reveal private chain-of-thought. Give concise reasoning and useful steps instead.

STYLE:
- Smart, calm, slightly casual human tutor.
- Concise for easy questions, detailed for difficult questions.
- Avoid customer-support language and repetitive canned greetings.
- Do not ask "How can I help?" after every greeting.
- Use Markdown naturally.

MATH:
- Use LaTeX.
- Display equations with $$ ... $$ and inline equations with $ ... $.
- Keep delimiters balanced.
- Explain important symbols and units.

FILES:
- Treat attached files as source material.
- Never invent facts that are not present in an attachment.
- Teach from extracted document content.
- For images, analyze only what is actually visible.
- Say what is missing if a file is unreadable/incomplete.
"""

app = FastAPI(title="Mentor.CaptainAI API", version="5.0.0")

# Extra allowed origins can be added without a code change: set
# EXTRA_CORS_ORIGINS="https://foo.com,https://bar.com" on Render.
_extra_origins = [o.strip() for o in os.getenv("EXTRA_CORS_ORIGINS", "").split(",") if o.strip()]

app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://127.0.0.1:5500",
        "http://localhost:5500",
        "https://mentor-captainai.89brats.workers.dev",
        *_extra_origins,
    ],
    # Covers any *.workers.dev / *.pages.dev preview or renamed Worker,
    # so redeploying the frontend under a new subdomain doesn't break CORS again.
    allow_origin_regex=r"https://.*\.(workers|pages)\.dev",
    allow_credentials=False,
    allow_methods=["GET", "POST", "OPTIONS"],
    allow_headers=["Content-Type"],
)

history_lock = threading.Lock()
conversation_history: dict[str, list[dict[str, str]]] = defaultdict(list)
upload_counts: dict[str, int] = defaultdict(int)


def ext(name: str) -> str:
    return Path(name).suffix.lower().lstrip(".")


def safe_name(name: str | None) -> str:
    return Path(name or "attachment").name[:180]


def read_text(data: bytes) -> str:
    return data.decode("utf-8", errors="replace")


def extract_document(name: str, data: bytes) -> str:
    e = ext(name)
    if e in {"txt","md","markdown","csv","json","py","js","ts","html","css","xml","yaml","yml","log","ini","toml","sql","java","c","cpp","h","hpp","jsx","tsx","sh","bat","ps1","env","rtf"}:
        return read_text(data)
    if e == "pdf":
        reader = PdfReader(io.BytesIO(data))
        return "\n\n".join(f"[PDF page {i}]\n{page.extract_text() or ''}" for i, page in enumerate(reader.pages, 1))
    if e == "docx":
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for n, table in enumerate(doc.tables, 1):
            parts.append(f"[DOCX table {n}]\n" + "\n".join(" | ".join(c.text.strip() for c in row.cells) for row in table.rows))
        return "\n\n".join(parts) or "[Empty DOCX]"
    if e == "pptx":
        prs = Presentation(io.BytesIO(data))
        parts = []
        for i, slide in enumerate(prs.slides, 1):
            texts = [shape.text.strip() for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()]
            parts.append(f"[Slide {i}]\n" + "\n".join(texts))
        return "\n\n".join(parts) or "[Empty PPTX]"
    if e == "xlsx":
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        parts = []
        for sheet in wb.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True):
                vals = ["" if v is None else str(v) for v in row]
                if any(vals):
                    rows.append(" | ".join(vals))
                if len(rows) >= 500:
                    break
            parts.append(f"[Worksheet: {sheet.title}]\n" + "\n".join(rows))
        return "\n\n".join(parts) or "[Empty XLSX]"
    return ""


def image_data_url(name: str, data: bytes) -> str:
    mime = {"png":"image/png","jpg":"image/jpeg","jpeg":"image/jpeg","webp":"image/webp","gif":"image/gif"}.get(ext(name), "application/octet-stream")
    return f"data:{mime};base64,{base64.b64encode(data).decode('ascii')}"


def get_history(session_id: str) -> list[dict[str, str]]:
    with history_lock:
        return list(conversation_history.get(session_id, [])[-MAX_HISTORY_MESSAGES:])


def save_turn(session_id: str, user: str, assistant: str) -> None:
    with history_lock:
        conversation_history[session_id].extend([
            {"role":"user","content":user},
            {"role":"assistant","content":assistant},
        ])
        conversation_history[session_id] = conversation_history[session_id][-MAX_HISTORY_MESSAGES:]


def api_error(response: requests.Response) -> str:
    try:
        body = response.json()
        return body.get("error", {}).get("message") or body.get("message") or response.text or f"HTTP {response.status_code}"
    except ValueError:
        return response.text or f"HTTP {response.status_code}"


def call_groq(session_id: str, user_text: str, docs: str, images: list[dict[str, Any]]) -> tuple[str, str]:
    if not GROQ_API_KEY:
        raise RuntimeError("GROQ_API_KEY is not configured on the server.")

    prompt = user_text.strip() or "Please inspect the attached file(s) and teach me what they contain."
    if docs.strip():
        prompt += "\n\nATTACHED SOURCE MATERIAL:\n" + docs[:MAX_DOCUMENT_CHARS]

    messages: list[dict[str, Any]] = [{"role":"system","content":SYSTEM_PROMPT}]
    messages.extend(get_history(session_id))

    model = VISION_MODEL if images else TEXT_MODEL
    if images:
        content: list[dict[str, Any]] = [{"type":"text","text":prompt}]
        content.extend(images)
        messages.append({"role":"user","content":content})
    else:
        messages.append({"role":"user","content":prompt})

    payload: dict[str, Any] = {
        "model": model,
        "messages": messages,
        "temperature": 0.7,
        "max_tokens": 1800,
    }
    if model.startswith("openai/gpt-oss"):
        payload["reasoning_effort"] = "medium"

    response = requests.post(
        GROQ_URL,
        headers={"Authorization": f"Bearer {GROQ_API_KEY}", "Content-Type":"application/json"},
        json=payload,
        timeout=120,
    )
    if response.status_code != 200:
        raise RuntimeError(api_error(response))
    try:
        answer = response.json()["choices"][0]["message"]["content"]
    except (KeyError, IndexError, TypeError, ValueError) as exc:
        raise RuntimeError("Groq returned an unexpected response.") from exc
    return str(answer).strip(), model


@app.get("/")
async def root():
    return {"status":"online","service":"Mentor.CaptainAI","version":"5.0.0","message":"Backend is running."}


@app.get("/health")
async def health():
    return {"status":"ok","service":"Mentor.CaptainAI"}


@app.post("/chat")
async def chat(request: Request):
    """Accept both the old JSON contract and the new multipart file-upload contract."""
    content_type = request.headers.get("content-type", "").lower()
    text = ""
    session_id = "default"
    files: list[UploadFile] = []

    if content_type.startswith("multipart/form-data"):
        form = await request.form()
        text = str(form.get("text") or "")
        session_id = str(form.get("session_id") or "default")
        for value in form.getlist("files"):
            if isinstance(value, UploadFile):
                files.append(value)
    elif content_type.startswith("application/json"):
        body = await request.json()
        text = str(body.get("text") or "")
        session_id = str(body.get("session_id") or "default")
    else:
        raise HTTPException(415, "Use application/json or multipart/form-data.")

    text = text.strip()[:5000]
    session_id = session_id.strip()[:128] or "default"
    current = upload_counts[session_id]

    if current + len(files) > MAX_FILES_PER_SESSION:
        raise HTTPException(429, f"This chat has reached the {MAX_FILES_PER_SESSION}-file upload limit.")
    if not text and not files:
        raise HTTPException(422, "Message or attachment required.")

    docs: list[str] = []
    images: list[dict[str, Any]] = []
    names: list[str] = []

    for upload in files:
        name = safe_name(upload.filename)
        e = ext(name)
        if e not in ALLOWED_EXTENSIONS:
            raise HTTPException(415, f"Unsupported file type: {name}")
        data = await upload.read()
        if len(data) > MAX_FILE_SIZE:
            raise HTTPException(413, f"{name} exceeds the 10 MB file limit.")
        names.append(name)
        if e in {"png","jpg","jpeg","webp","gif"}:
            if len(images) >= MAX_IMAGES_PER_REQUEST:
                raise HTTPException(400, f"Only {MAX_IMAGES_PER_REQUEST} images may be attached to one message.")
            images.append({"type":"image_url","image_url":{"url":image_data_url(name, data)}})
        else:
            docs.append(f"--- FILE: {name} ---\n{extract_document(name, data)}\n--- END FILE: {name} ---")

    history_user = text or "Please inspect the attached file(s)."
    if names:
        history_user += "\n\nAttached files: " + ", ".join(names)

    try:
        answer, model_used = call_groq(session_id, text, "\n\n".join(docs), images)
    except requests.Timeout as exc:
        raise HTTPException(504, "The AI service timed out. Please try again.") from exc
    except requests.RequestException as exc:
        raise HTTPException(502, "The AI service could not be reached.") from exc
    except RuntimeError as exc:
        message = str(exc)
        if "blocked at the project level" in message.lower():
            message += " Enable the configured VISION_MODEL in your Groq project, or set VISION_MODEL to another permitted vision model."
        raise HTTPException(502, message) from exc

    save_turn(session_id, history_user, answer)
    upload_counts[session_id] = current + len(files)

    return {
        "response": answer,
        "uploaded_files": len(files),
        "uploads_used": upload_counts[session_id],
        "uploads_remaining": MAX_FILES_PER_SESSION - upload_counts[session_id],
        "model_used": model_used,
    }
